#!/usr/bin/env python3
"""Build a one-page resume by filling the immutable bundled DOCX template."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Iterable

from docx import Document

from validate_resume_content import (
    DEFAULT_PROFILE,
    DEFAULT_PROJECTS,
    load_json,
    validate,
)


SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATES = {
    "embedded": (
        SKILL_DIR / "assets" / "embedded-master-20260904.docx",
        "FD387ED647C59727CF0E6B99B0839377FF8015C85592A4AD587013D394A3433B",
    ),
    "general": (
        SKILL_DIR / "assets" / "general-master.docx",
        "E271C561F43CC5D16CE32F3B555F349009EF1F2165C1EE39CBD57538C78172D4",
    ),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def replace_runs(paragraph, segments: Iterable[tuple[str, int]]) -> None:
    old_runs = list(paragraph.runs)
    if not old_runs:
        raise ValueError("template paragraph has no style-bearing run")
    styles = [deepcopy(run._r.rPr) if run._r.rPr is not None else None for run in old_runs]
    for run in old_runs:
        paragraph._p.remove(run._r)
    for text, style_index in segments:
        new_run = paragraph.add_run(text)
        style = styles[min(style_index, len(styles) - 1)]
        if style is not None:
            new_run._r.insert(0, deepcopy(style))


def replace_labelled(paragraph, label: str, text: str, bullet: bool) -> None:
    prefix = f"• {label}：" if bullet else f"{label}："
    replace_runs(paragraph, [(prefix, 0), (text, 1)])


def replace_title_table(table, name: str, role: str, period: str) -> None:
    left = table.cell(0, 0).paragraphs[0]
    right = table.cell(0, 1).paragraphs[0]
    replace_runs(left, [(name, 0), (f"｜{role}", 1)])
    replace_runs(right, [(period, 0)])


def flatten_project_bullets(projects: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [bullet for project in projects for bullet in project["bullets"]]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--content", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--profile", type=Path, default=DEFAULT_PROFILE)
    parser.add_argument("--projects", type=Path, default=DEFAULT_PROJECTS)
    args = parser.parse_args()

    if args.output.exists():
        print(f"ERROR: output already exists: {args.output}", file=sys.stderr)
        return 2

    try:
        payload = load_json(args.content)
        profile = load_json(args.profile)
        project_evidence = load_json(args.projects)
    except (OSError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    errors, warnings = validate(payload, profile, project_evidence)
    for warning in warnings:
        print(f"WARNING: {warning}")
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 2

    template_name = payload["template"]
    template_path, expected_hash = TEMPLATES[template_name]
    actual_hash = sha256(template_path)
    if actual_hash != expected_hash:
        print(
            f"ERROR: template hash mismatch for {template_path}: {actual_hash} != {expected_hash}",
            file=sys.stderr,
        )
        return 2

    document = Document(template_path)
    if len(document.paragraphs) != 27 or len(document.tables) != 7:
        print(
            f"ERROR: unsupported template structure: {len(document.paragraphs)} paragraphs, "
            f"{len(document.tables)} tables",
            file=sys.stderr,
        )
        return 2

    header = payload["header"]
    header_cell = document.tables[0].cell(0, 0)
    replace_runs(
        header_cell.paragraphs[0],
        [(profile["person"]["name"], 0), (f"   求职意向：{header['target']}", 1)],
    )
    replace_runs(
        header_cell.paragraphs[1],
        [(
            f"电话：{profile['person']['phone']}   邮箱：{profile['person']['email']}   "
            f"{profile['person']['graduation_year']}届",
            0,
        )],
    )
    replace_runs(header_cell.paragraphs[2], [(header["tags"], 0)])

    for paragraph_index, skill in zip((2, 3, 4), payload["skills"]):
        replace_labelled(document.paragraphs[paragraph_index], skill["label"], skill["text"], False)

    for table_index, project in zip((2, 3, 4), payload["projects"]):
        replace_title_table(document.tables[table_index], project["name"], project["role"], project["period"])

    bullet_slots = (6, 7, 8, 9, 10, 11, 12, 13)
    for paragraph_index, bullet in zip(bullet_slots, flatten_project_bullets(payload["projects"])):
        replace_labelled(document.paragraphs[paragraph_index], bullet["label"], bullet["text"], True)

    enterprise = payload["enterprise"]
    replace_title_table(
        document.tables[5], enterprise["name"], enterprise["role"], enterprise["period"]
    )
    for paragraph_index, bullet in zip((15, 16), enterprise["bullets"]):
        replace_labelled(document.paragraphs[paragraph_index], bullet["label"], bullet["text"], True)

    campus = payload["campus"]
    replace_title_table(document.tables[6], campus["name"], campus["role"], campus["period"])
    for paragraph_index, bullet in zip((18, 19), campus["bullets"]):
        replace_labelled(document.paragraphs[paragraph_index], bullet["label"], bullet["text"], True)

    for paragraph_index, item in zip((21, 22, 23, 24), payload["competitions"]):
        replace_labelled(document.paragraphs[paragraph_index], item["label"], item["text"], False)

    self_evaluation = payload["self_evaluation"]
    replace_labelled(
        document.paragraphs[26], self_evaluation["label"], self_evaluation["text"], False
    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = f"{profile['person']['name']} - {payload['company']} - {payload['role']}"
    document.save(args.output)

    try:
        with zipfile.ZipFile(args.output) as archive:
            bad_entry = archive.testzip()
    except zipfile.BadZipFile as exc:
        print(f"ERROR: generated DOCX is invalid: {exc}", file=sys.stderr)
        return 2
    if bad_entry is not None:
        print(f"ERROR: corrupt DOCX member: {bad_entry}", file=sys.stderr)
        return 2

    print(f"OK: generated {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
