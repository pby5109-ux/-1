# -*- coding: utf-8 -*-
"""青岛乾程科技-嵌入式软件工程师 定制（组合01：实习有、校园无）。
底版=磁盘-01.docx；内容按用户 09-28 08:40 自导出 -01.pdf 口径（权威）：
LoRa槽位简介/职责恢复、瞄准简介删"并开启激光"、模型训练尾部"的"、融合定位=用户-01改写版
（AI视觉/几何识别双策略切换）、功耗优化缩句修句号孤行；技能通信与协议加TCP/UDP、HTTP
（用户实验级证据+JD要求TCP/IP；Modbus/CAN/Linux无证据不写）；三项目块灰色分界线。"""
import copy, hashlib, zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"D:\APP\project_practice\codex-project\简历\简历修改版本\模板简历（随时跟新）\模板简历\彭博裕-嵌入式软件工程师-01.docx")
EXPECTED = "B4819DDF5F12164D"
OUT = Path(r"D:\APP\project_practice\codex-project\简历\中间文件\青岛乾程科技-20260929\彭博裕-嵌入式软件工程师-01.docx")

AIM_INTRO = "面向循迹小车行驶中的激光自动瞄准，结合K230 KPU高速推理单元，实现AI目标检测与几何视觉，完成矩形靶标识别、靶心定位及瞄准偏差反馈。不同环境下目标识别率达98%；整机实测启动后2秒内完成瞄准，打靶误差＜1.5cm，小车稳定行驶时误差＜2cm。"
P2_模型训练 = "为解决各种环境对识别精度影响，自采多环境靶标图像，通过Make Sense标注并借助AI校验标注质量；在AI Cube完成目标检测模型训练与板端部署，实现不同环境下的矩形靶标的准确识别。"
P2_融合定位 = "将AI检测框与几何四角统一到同帧原图坐标并匹配，优先通过四角对角线交点求心；几何失效但AI有效时通过用户自动切换策略，实现AI视觉和几何识别灵活切换，大幅提高运动过程中的识别精度和抗风险能力。"
LORA_INTRO = "基于STM32F103C8T6，完成集环境数据采集、周期无线上报与网关异常记录于一体的低功耗监测节点，实现整机休眠电流实测约54μA；按1100mAh电池容量及既定上报工况，预留裕量后估算续航约8个月。"
LORA_DUTY = "负责方案选型、软件开发与整板调试。"
P3_功耗优化 = "结合Stop模式、代码优化、供电门控与休眠控制，抑制GPIO漏电并隔离断电外设信号，将整板待机电流控制在µA级。"
SKILL_通信与协议 = "熟悉UART、I²C、SPI、TCP/UDP、HTTP等常用通信协议与自定义帧协议设计，使用过通信、采集与定时控制类外设。"
INTERNSHIP = [
    ("产品认知", "结合智能电表及采集终端PCBA，梳理电源、采样、计量、通信与防护模块，理解三相四线计量及主要器件连接关系。"),
    ("开发板实践", "参与电能表开发板操作及测试，结合培训梳理工况事件、远程费控、主动上报的输入条件、软件处理与设备响应。"),
    ("通信与排查", "学习RS485、电力线载波及耦合原理，整理“供电接线—信号输入—通信参数—软件状态”分层排查思路，辅助问题定位。"),
]
BORDER_LABELS = ("测量与调试", "稳健反馈", "协议开发")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for c in iter(lambda: f.read(1 << 20), b''):
            h.update(c)
    return h.hexdigest().upper()


def replace_runs(paragraph, segments):
    old = list(paragraph.runs)
    if not old:
        raise ValueError('空段落')
    styles = [copy.deepcopy(r._r.rPr) if r._r.rPr is not None else None for r in old]
    for r in old:
        paragraph._p.remove(r._r)
    for text, idx in segments:
        run = paragraph.add_run(text)
        st = styles[min(idx, len(styles) - 1)]
        if st is not None:
            run._r.insert(0, copy.deepcopy(st))


def labelled(paragraph, label, text, bullet=False):
    prefix = f"• {label}：" if bullet else f"{label}："
    replace_runs(paragraph, [(prefix, 0), (text, 1)])


def add_gray_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '808080')
    pBdr.append(bottom)
    spacing = pPr.find(qn('w:spacing'))
    if spacing is not None:
        spacing.addprevious(pBdr)
    else:
        pPr.append(pBdr)


def main():
    if not sha256(BASE).startswith(EXPECTED):
        print('ERROR: 基准哈希不符'); return 2
    doc = Document(BASE)

    done = []
    for p in doc.paragraphs:
        t = p.text.strip()
        core = t[2:] if t.startswith("• ") else t
        label = core.split("：", 1)[0]
        if t.startswith("项目简介：面向小车行驶过程中的激光自动瞄准"):
            labelled(p, "项目简介", LORA_INTRO); done.append("LoRa简介")
        elif t == "项目职责：负责视觉算法、灰度采集与OLED/按键交互。":
            labelled(p, "项目职责", LORA_DUTY); done.append("LoRa职责")
        elif t.startswith("项目简介：面向循迹小车行驶中的激光自动瞄准"):
            labelled(p, "项目简介", AIM_INTRO); done.append("瞄准简介")
        elif label == "模型训练":
            labelled(p, label, P2_模型训练, bullet=True); done.append("模型训练")
        elif label == "融合定位":
            labelled(p, label, P2_融合定位, bullet=True); done.append("融合定位(PDF新版)")
        elif label == "功耗优化":
            labelled(p, label, P3_功耗优化, bullet=True); done.append("功耗优化缩句")
        elif label == "通信与协议":
            labelled(p, label, SKILL_通信与协议); done.append("通信与协议+TCP/UDP/HTTP")
        elif label in BORDER_LABELS:
            add_gray_bottom_border(p); done.append(f"分界线-{label}")
        elif label == "实习经历条目核对":
            pass

    need = {"LoRa简介", "LoRa职责", "瞄准简介", "模型训练", "融合定位(PDF新版)", "功耗优化缩句",
            "通信与协议+TCP/UDP/HTTP", "分界线-测量与调试", "分界线-稳健反馈", "分界线-协议开发"}
    if set(done) != need:
        print(f'ERROR: 补丁命中不全：缺 {need - set(done)} 多 {set(done) - need}'); return 2

    # 实习三条核对：与用户 -01.pdf 口径不一致则替换
    intern_hit = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        core = t[2:] if t.startswith("• ") else t
        label = core.split("：", 1)[0]
        for want_label, want_text in INTERNSHIP:
            if label == want_label and core[len(label) + 1:] != want_text:
                labelled(p, want_label, want_text, bullet=True)
                done.append(f"实习-{want_label}-已按PDF替换")
                intern_hit += 1
            elif label == want_label:
                intern_hit += 1
    if intern_hit != 3:
        print(f'ERROR: 实习三条仅命中 {intern_hit}/3'); return 2

    OUT.parent.mkdir(parents=True, exist_ok=True)
    staged = OUT.with_name(OUT.stem + '.staged.docx')
    doc.save(staged)
    with zipfile.ZipFile(staged) as e:
        edited = e.read('word/document.xml')
    with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as o, zipfile.ZipFile(BASE) as z:
        for item in z.infolist():
            o.writestr(item, edited if item.filename == 'word/document.xml' else z.read(item.filename))
    staged.unlink()

    diff = []
    with zipfile.ZipFile(BASE) as a, zipfile.ZipFile(OUT) as b:
        assert b.testzip() is None
        for name in a.namelist():
            if hashlib.sha256(a.read(name)).hexdigest() != hashlib.sha256(b.read(name)).hexdigest():
                diff.append(name)
    print("已改：", "、".join(done)); print("差异部件：", diff)
    if diff != ['word/document.xml']:
        print('ERROR: 差异部件异常'); return 2
    print('OK:', OUT)
    return 0


raise SystemExit(main())
