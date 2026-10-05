# -*- coding: utf-8 -*-
"""青岛乾程版 v2：基于更新后的 -01 母版（新实习三条+口径补丁+分界线已在母版中），
仅追加乾程专属技能行：通信与协议加 TCP/UDP、HTTP（JD要求TCP/IP，实验级证据）。"""
import copy, hashlib, zipfile
from pathlib import Path
from docx import Document

BASE = Path(r"D:\APP\project_practice\codex-project\简历\简历修改版本\模板简历（随时跟新）\模板简历\彭博裕-嵌入式软件工程师-01.docx")
EXPECTED = "5992566855149656"
OUT = Path(r"D:\APP\project_practice\codex-project\简历\中间文件\青岛乾程科技-20260929\彭博裕-嵌入式软件工程师-01.docx")

SKILL_通信与协议 = "熟悉UART、I²C、SPI、TCP/UDP、HTTP等常用通信协议与自定义帧协议设计，使用过通信、采集与定时控制类外设。"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for c in iter(lambda: f.read(1 << 20), b''):
            h.update(c)
    return h.hexdigest().upper()


def replace_runs(paragraph, segments):
    old = list(paragraph.runs)
    if not old:
        raise ValueError('空段落')
    styles = [copy.deepcopy(r._r.rPr) if r._r.rPr is not None else None for r in old]
    for r in old:
        paragraph._p.remove(r._r)
    for text, idx in segments:
        run = paragraph.add_run(text)
        st = styles[min(idx, len(styles) - 1)]
        if st is not None:
            run._r.insert(0, copy.deepcopy(st))


def main():
    if not sha256(BASE).startswith(EXPECTED):
        print(f'ERROR: 基准哈希不符：{sha256(BASE)[:16]}'); return 2
    doc = Document(BASE)

    hit = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.split("：", 1)[0] == "通信与协议":
            replace_runs(p, [("通信与协议：", 0), (SKILL_通信与协议, 1)])
            hit += 1
    if hit != 1:
        print(f'ERROR: 通信与协议命中 {hit}/1'); return 2

    OUT.parent.mkdir(parents=True, exist_ok=True)
    staged = OUT.with_name(OUT.stem + '.staged.docx')
    doc.save(staged)
    with zipfile.ZipFile(staged) as e:
        edited = e.read('word/document.xml')
    with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as o, zipfile.ZipFile(BASE) as z:
        for item in z.infolist():
            o.writestr(item, edited if item.filename == 'word/document.xml' else z.read(item.filename))
    staged.unlink()

    diff = []
    with zipfile.ZipFile(BASE) as a, zipfile.ZipFile(OUT) as b:
        assert b.testzip() is None
        for name in a.namelist():
            if hashlib.sha256(a.read(name)).hexdigest() != hashlib.sha256(b.read(name)).hexdigest():
                diff.append(name)
    print("差异部件：", diff)
    if diff != ['word/document.xml']:
        print('ERROR: 差异部件异常'); return 2
    print('OK:', OUT)
    return 0


raise SystemExit(main())
