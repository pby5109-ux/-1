# -*- coding: utf-8 -*-
"""无人机配套MCU嵌入式（深圳）定制：底版=磁盘-10.docx，内容按用户09-28 08:40 PDF口径（权威）：
LoRa简介/职责恢复54µA文案、瞄准简介删"并开启激光"、融合定位=用户PDF原文（用户确认有意，非误粘贴）、
功耗优化缩句修"级。"孤行；求职意向改JD岗位名；三项目块尾加灰色分界线。技能行保持母版（无IAR/TCP-UDP-HTTP）。"""
import copy, hashlib, zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

BASE = Path(r"D:\APP\project_practice\codex-project\简历\简历修改版本\模板简历（随时跟新）\模板简历\彭博裕-嵌入式软件工程师-10.docx")
EXPECTED = "E911EF38DCA4D0C6"
OUT = Path(r"D:\APP\project_practice\codex-project\简历\中间文件\无人机配套MCU-20260928\彭博裕-MCU嵌入式软件开发工程师-10.docx")

TARGET = "MCU嵌入式软件开发工程师"
LORA_INTRO = "基于STM32F103C8T6，完成集环境数据采集、周期无线上报与网关异常记录于一体的低功耗监测节点，实现整机休眠电流实测约54µA；按1100mAh电池容量及既定上报工况，预留裕量后估算续航约8个月。"
LORA_DUTY = "负责方案选型、软件开发与整板调试。"
AIM_INTRO = "面向循迹小车行驶中的激光自动瞄准，结合K230 KPU高速推理单元，实现AI目标检测与几何视觉，完成矩形靶标识别、靶心定位及瞄准偏差反馈。不同环境下目标识别率达98%；整机实测启动后2秒内完成瞄准，打靶误差＜1.5cm，小车稳定行驶时误差＜2cm。"
FUSION_PDF = "为解决各种环境对识别精度影响，自采多环境靶标图像，通过Make Sense标注并借助AI校验标注质量；在AI Cube完成目标检测模型训练与板端部署，实现不同环境下的矩形靶标的准确识别。"
P3_功耗优化 = "结合Stop模式、代码优化、供电门控与休眠控制，抑制GPIO漏电并隔离断电外设信号，将整板待机电流控制在µA级。"
BORDER_LABELS = ("测量与调试", "稳健反馈", "协议开发")   # 三个项目块末条：灰色分界线


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for c in iter(lambda: f.read(1 << 20), b''):
            h.update(c)
    return h.hexdigest().upper()


def replace_runs(paragraph, segments):
    old = list(paragraph.runs)
    if not old:
        raise ValueError('空段落')
    styles = [copy.deepcopy(r._r.rPr) if r._r.rPr is not None else None for r in old]
    for r in old:
        paragraph._p.remove(r._r)
    for text, idx in segments:
        run = paragraph.add_run(text)
        st = styles[min(idx, len(styles) - 1)]
        if st is not None:
            run._r.insert(0, copy.deepcopy(st))


def labelled(paragraph, label, text, bullet=False):
    prefix = f"• {label}：" if bullet else f"{label}："
    replace_runs(paragraph, [(prefix, 0), (text, 1)])


def add_gray_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')        # 0.75pt
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '808080')
    pBdr.append(bottom)
    spacing = pPr.find(qn('w:spacing'))
    if spacing is not None:
        spacing.addprevious(pBdr)      # schema 顺序：pBdr 在 spacing 之前
    else:
        pPr.append(pBdr)


def main():
    if not sha256(BASE).startswith(EXPECTED):
        print('ERROR: 基准哈希不符'); return 2
    doc = Document(BASE)

    done = []
    for p in doc.paragraphs:
        t = p.text.strip()
        core = t[2:] if t.startswith("• ") else t
        label = core.split("：", 1)[0]
        if t.startswith("项目简介：面向小车行驶过程中的激光自动瞄准"):
            labelled(p, "项目简介", LORA_INTRO); done.append("LoRa简介")
        elif t == "项目职责：负责视觉算法、灰度采集与OLED/按键交互。":
            labelled(p, "项目职责", LORA_DUTY); done.append("LoRa职责")
        elif t.startswith("项目简介：面向循迹小车行驶中的激光自动瞄准"):
            labelled(p, "项目简介", AIM_INTRO); done.append("瞄准简介")
        elif label == "融合定位":
            labelled(p, label, FUSION_PDF, bullet=True); done.append("融合定位(PDF口径)")
        elif label == "功耗优化":
            labelled(p, label, P3_功耗优化, bullet=True); done.append("功耗优化缩句")
        elif label in BORDER_LABELS:
            add_gray_bottom_border(p); done.append(f"分界线-{label}")
    need = {"LoRa简介", "LoRa职责", "瞄准简介", "融合定位(PDF口径)", "功耗优化缩句",
            "分界线-测量与调试", "分界线-稳健反馈", "分界线-协议开发"}
    if set(done) != need:
        print(f'ERROR: 补丁命中不全：缺 {need - set(done)} 多 {set(done) - need}'); return 2

    # 求职意向（表头 T0）
    cell = doc.tables[0].cell(0, 0)
    hit = False
    for p in cell.paragraphs:
        for r in p.runs:
            if "求职意向" in r.text:
                r.text = f"求职意向：{TARGET}"; hit = True
    if not hit:
        for p in cell.paragraphs:
            if "求职意向" in p.text:
                full = p.text
                replace_runs(p, [(f"求职意向：{TARGET}", 0)] if len(p.runs) <= 1 else
                             [(s.replace("嵌入式软件工程师", TARGET), i) for i, s in
                              enumerate([r.text for r in p.runs])])
                hit = True; break
    if not hit:
        print('ERROR: 未找到求职意向'); return 2
    done.append("求职意向")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    staged = OUT.with_name(OUT.stem + '.staged.docx')
    doc.save(staged)
    with zipfile.ZipFile(staged) as e:
        edited = e.read('word/document.xml')
    with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as o, zipfile.ZipFile(BASE) as z:
        for item in z.infolist():
            o.writestr(item, edited if item.filename == 'word/document.xml' else z.read(item.filename))
    staged.unlink()

    diff = []
    with zipfile.ZipFile(BASE) as a, zipfile.ZipFile(OUT) as b:
        assert b.testzip() is None
        for name in a.namelist():
            if hashlib.sha256(a.read(name)).hexdigest() != hashlib.sha256(b.read(name)).hexdigest():
                diff.append(name)
    print("已改：", "、".join(done)); print("差异部件：", diff)
    if diff != ['word/document.xml']:
        print('ERROR: 差异部件异常'); return 2
    print('OK:', OUT)
    return 0


raise SystemExit(main())
