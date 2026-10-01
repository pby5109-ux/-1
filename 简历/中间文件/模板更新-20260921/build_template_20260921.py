#!/usr/bin/env python3
"""删除实习经历栏并把整体字号提一级，输出 1 页 PDF 用 DOCX 中间件。

基准：简历/简历修改版本/模板简历（随时跟新）/彭博裕-嵌入式软件工程师-含实习无校园.docx
只替换 word/document.xml，其余 ZIP 部件与基准逐字节一致。
"""
from __future__ import annotations

import copy
import hashlib
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

BASE = Path(
    r"D:\APP\project_practice\codex-project\简历\简历修改版本\模板简历（随时跟新）"
    r"\彭博裕-嵌入式软件工程师-含实习无校园.docx"
)
EXPECTED_HASH = "C42B86ED97E858E116FCD5524ACFDADA49CE5E7E97C41AEB7F2B3747178346F2"

# 实习经历栏在 body 中的元素下标（含栏目标题、机构标题表、三条正文）
INTERNSHIP_INDEXES = (26, 27, 28, 29, 30)
# 字号映射（半磅）：正文 8→9pt，标题类各 +0.5pt
SIZE_MAP = {"16": "18", "18": "19", "19": "20", "20": "21"}
LINE_TWIPS = "226"          # 行距 0.942（基准 230）
HEADING_BEFORE_TWIPS = "20"  # 栏目标题段前 1pt（基准 60=3pt）


# 用户 2026-09-20 版（仅存在于 PDF，本处按 PDF 文本层重建）相对 09-18 DOCX 的正文差异
TEXT_OVERRIDES = {
    "项目简介#1": "集成示波器、数字万用表、可调直流电源和信号发生器，支持波形实时显示及频率、幅度测量，频率测量误差1 Hz，电压与电阻测量误差1%；支持直流电压稳定输出；信号发生器输出正弦、方波等，频率档位覆盖50 Hz～50 kHz。",
    "并发优化": "针对采集与显示共享缓冲区的读写冲突，设计双缓冲与状态机协同机制，管理波形数据的读写权限和切换时机，保证单帧绘制期间数据一致，避免新旧帧混用，并允许新帧准备与当前帧绘制并行。",
    "测量与调试": "基于内部参考电压估算VDDA，结合GPIO档位识别、分压与偏置补偿完成多量程换算，提升测量准确性；优化采样率与波表长度配置、定时参数计算及更新时序，修正低频输出偏差，实现频率切换及时生效。",
    "系统开发": "分层封装传感器驱动与采集逻辑，设计RTC周期采集与按键唤醒双事件机制，分别维护采样与显示期限，避免临时查看打乱上报节拍；支持无网关条件下现场查看数据，并通过有效性检查标识采集异常，保障不同情况下数据的有效观测。",
    "功耗优化": "结合Stop模式、代码系统优化、外设供电门控与休眠控制，抑制GPIO漏电并隔离断电外设信号，将整板待机电流控制在µA级。",
    "协议开发": "设计CRC校验、序号匹配及ACK超时有限重试机制，实现上报确认与传输异常处理；开发PC网关处理半包、粘包及坏帧，按节点与序号去重且对重传仍回ACK，完成数据解析及带时间的故障日志记录，支持采集异常追溯。",
    "光照适配": "基于Otsu算法自动估计灰度阈值，通过按键触发重估、帧间保持阈值，实现不同光环境下的阈值校准；对无灰度对比的画面跳过识别并重试，避免无效画面参与阈值估计。",
    "AI工具": "熟练使用Codex、Claude Code、DeepSeek Harness、ZCode搭建工作流并实现快速产品设计与开发。",
}


def apply_overrides(document, overrides) -> int:
    """按「标签」定位段落并替换正文，保留原有加粗标签＋正文的两段式格式。"""
    seen: dict[str, int] = {}
    applied = 0
    for paragraph in document.paragraphs:
        current = paragraph.text
        if "：" not in current:
            continue
        bullet = current.lstrip().startswith("•")
        body = current.lstrip().lstrip("•").strip()
        label = body.split("：", 1)[0].strip()
        seen[label] = seen.get(label, 0) + 1
        key = label if label not in ("项目简介", "项目职责") else f"{label}#{seen[label]}"
        if key not in overrides:
            continue
        prefix = f"• {label}：" if bullet else f"{label}："
        replace_runs(paragraph, [(prefix, 0), (overrides[key], 1)])
        applied += 1
    return applied


def tighten_layout(document, line_twips: str = "222", heading_before_twips: str = "20") -> tuple[int, int]:
    """字号放大后收紧行距与栏目标题段前距，保证仍为一页。"""
    headings = {"教育背景", "竞赛与荣誉", "项目经历", "实习经历", "专业技能"}
    line_changed = before_changed = 0
    for paragraph in document.paragraphs:
        pPr = paragraph._p.find(qn("w:pPr"))
        if pPr is None:
            continue
        spacing = pPr.find(qn("w:spacing"))
        if spacing is None:
            continue
        if spacing.get(qn("w:line")):
            spacing.set(qn("w:line"), line_twips)
            line_changed += 1
        if paragraph.text.strip() in headings and spacing.get(qn("w:before")) not in (None, "0"):
            spacing.set(qn("w:before"), heading_before_twips)
            before_changed += 1
    return line_changed, before_changed


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def replace_runs(paragraph, segments) -> None:
    old_runs = list(paragraph.runs)
    if not old_runs:
        raise ValueError("段落没有可继承格式的 run")
    styles = [copy.deepcopy(run._r.rPr) if run._r.rPr is not None else None for run in old_runs]
    for run in old_runs:
        paragraph._p.remove(run._r)
    for text, style_index in segments:
        new_run = paragraph.add_run(text)
        style = styles[min(style_index, len(styles) - 1)]
        if style is not None:
            new_run._r.insert(0, copy.deepcopy(style))


def main() -> int:
    output = Path(sys.argv[1])
    if output.exists():
        print(f"ERROR: 输出已存在：{output}", file=sys.stderr)
        return 2
    actual = sha256(BASE)
    if actual != EXPECTED_HASH:
        print(f"ERROR: 基准哈希不符\n实际 {actual}\n期望 {EXPECTED_HASH}", file=sys.stderr)
        return 2
    print(f"基准哈希核对通过：{actual[:16]}…")

    doc = Document(BASE)
    body = doc.element.body
    children = list(body)

    applied = apply_overrides(doc, TEXT_OVERRIDES) if TEXT_OVERRIDES else 0
    print(f"按 09-20 版正文更新段落：{applied} 处")

    removed = 0
    for index in INTERNSHIP_INDEXES:
        element = children[index]
        if element.getparent() is not None:
            element.getparent().remove(element)
            removed += 1
    print(f"已删除实习经历栏（{removed} 个元素：栏目标题／机构标题行／三条正文）")

    changed = 0
    for run in body.iter(qn("w:r")):
        rPr = run.find(qn("w:rPr"))
        if rPr is None:
            continue
        for tag in ("w:sz", "w:szCs"):
            size = rPr.find(qn(tag))
            if size is None:
                continue
            value = size.get(qn("w:val"))
            if value in SIZE_MAP:
                size.set(qn("w:val"), SIZE_MAP[value])
                changed += 1
    print(f"字号提升：{changed} 处（映射 {SIZE_MAP}）")

    line_n, head_n = tighten_layout(doc, LINE_TWIPS, HEADING_BEFORE_TWIPS)
    print(f"版面收紧：行距 {line_n} 处 → {LINE_TWIPS}，栏目标题段前距 {head_n} 处 → {HEADING_BEFORE_TWIPS}")


    parent = output.parent
    parent.mkdir(parents=True, exist_ok=True)
    staged = parent / (output.stem + ".staged.docx")
    doc.save(staged)

    with zipfile.ZipFile(BASE) as zin, zipfile.ZipFile(staged) as zedit:
        edited = zedit.read("word/document.xml")
    with zipfile.ZipFile(BASE) as zin, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = edited if item.filename == "word/document.xml" else zin.read(item.filename)
            zout.writestr(item, data)
    staged.unlink()

    with zipfile.ZipFile(output) as archive:
        bad = archive.testzip()
    if bad is not None:
        print(f"ERROR: DOCX 损坏成员：{bad}", file=sys.stderr)
        return 2
    print(f"OK: 已生成 {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
