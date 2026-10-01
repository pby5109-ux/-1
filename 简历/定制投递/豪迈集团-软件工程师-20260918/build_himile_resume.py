#!/usr/bin/env python3
"""豪迈集团 · 软件工程师 定制简历生成。

基准：简历/简历修改版本/模板简历（随时跟新）/彭博裕-嵌入式软件工程师-含实习与校园.docx
差量：求职意向改软件工程师；项目组合换成 仪器＋节点＋农场（换出瞄准系统）；
      仪器与节点各换三个关键词；删除校园经历整栏。其余正文、样式、媒体保持不变。
"""

from __future__ import annotations

import copy
import hashlib
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

BASE = Path(
    r"D:\APP\project_practice\codex-project\简历\简历修改版本\模板简历（随时跟新）"
    r"\彭博裕-嵌入式软件工程师-含实习与校园.docx"
)
EXPECTED_HASH = "E6E1227599E3A226CE0A1EA80887E72228D3DD18A7FCBC637F9AE404A6206236"

TARGET = "软件工程师"

KEYWORDS = {
    "instrument": "FreeRTOS · 软件分层 · 系统优化",
    "lora": "上位机网关 · 协议解析 · 低功耗管理",
    "farm": "多任务架构 · 状态解耦 · 可靠性优化",
}

FARM_NAME = "基于STM32与FreeRTOS的智慧农场环境监测与控制系统"
FARM_INTRO = (
    "基于STM32F103C8T6与FreeRTOS的环境监测与自动控制系统，采集空气温湿度、土壤湿度、"
    "光照及降雨状态，按安全阈值驱动水泵与风扇并推送蓝牙告警，形成异常检测、执行器联动与恢复通知的控制闭环。"
)
FARM_DUTY = "核心开发与系统优化。"
FARM_BULLETS = [
    (
        "任务架构",
        "基于STM32F103C8T6与FreeRTOS搭建环境监测和自动控制系统，将传感采集、按键/编码器、"
        "OLED显示和BLE通信划分为Sensor、Input、Screen、BLE四个任务，并以farmState/farmSafeRange"
        "解耦实时状态、界面显示与可配置安全阈值。",
    ),
    (
        "数据采集",
        "结合AHT20、BH1750及土壤/雨滴传感器采集空气温湿度、土壤温湿度、光照和降雨相对数据，"
        "使用ADC1扫描与循环DMA、ADC2连续转换及I²C完成多源读取；以互斥锁保护AHT20与OLED共享的I²C1总线。",
    ),
    (
        "控制告警",
        "基于阈值驱动水泵和PWM风扇，通过指针消息队列与UART DMA发送JSON告警；修复队列满时动态消息"
        "未释放造成的内存泄漏，将持续报警改为warning/recovered状态边沿，并为水泵加入5%滞回，"
        "减少重复消息与临界启停。",
    ),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def replace_runs(paragraph: Paragraph, segments) -> None:
    old_runs = list(paragraph.runs)
    if not old_runs:
        raise ValueError("段落没有可继承格式的 run")
    styles = [copy.deepcopy(run._r.rPr) if run._r.rPr is not None else None for run in old_runs]
    for run in old_runs:
        paragraph._p.remove(run._r)
    for text, style_index in segments:
        new_run = paragraph.add_run(text)
        style = styles[min(style_index, len(styles) - 1)]
        if style is not None:
            new_run._r.insert(0, copy.deepcopy(style))


def labelled(paragraph: Paragraph, label: str, text: str, bullet: bool = False) -> None:
    prefix = f"• {label}：" if bullet else f"{label}："
    replace_runs(paragraph, [(prefix, 0), (text, 1)])


def main() -> int:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("彭博裕-豪迈集团-软件工程师-20260918.docx")
    if output.exists():
        print(f"ERROR: 输出已存在：{output}", file=sys.stderr)
        return 2
    actual = sha256(BASE)
    if actual != EXPECTED_HASH:
        print(f"ERROR: 基准模板哈希不符，停止生成。\n实际 {actual}\n期望 {EXPECTED_HASH}", file=sys.stderr)
        return 2
    print(f"基准 SHA256 核对通过: {actual[:16]}…")

    doc = Document(BASE)
    body = doc.element.body
    children = list(body)
    paras = doc.paragraphs
    tables = doc.tables

    # 1) 求职意向
    header_cell = tables[0].cell(0, 0)
    for run in header_cell.paragraphs[0].runs:
        if "嵌入式软件工程师" in run.text:
            run.text = run.text.replace("嵌入式软件工程师", TARGET)
    print("求职意向 ->", TARGET)

    # 2) 三个项目标题行右格关键词
    replace_runs(tables[2].cell(0, 1).paragraphs[0], [(KEYWORDS["instrument"], 0)])
    replace_runs(tables[3].cell(0, 1).paragraphs[0], [(KEYWORDS["lora"], 0)])
    print("关键词已替换：仪器 / 节点")

    # 3) 第三个项目块：瞄准系统 -> 智慧农场
    replace_runs(tables[4].cell(0, 0).paragraphs[0], [(FARM_NAME, 0)])
    replace_runs(tables[4].cell(0, 1).paragraphs[0], [(KEYWORDS["farm"], 0)])
    project_paragraphs = [Paragraph(children[i], doc) for i in (21, 22, 23, 24, 25)]
    replace_runs(project_paragraphs[0], [("项目简介：", 0), (FARM_INTRO, 1)])
    replace_runs(project_paragraphs[1], [("项目职责：", 0), (FARM_DUTY, 1)])
    for (label, text), paragraph in zip(FARM_BULLETS, project_paragraphs[2:]):
        labelled(paragraph, label, text, bullet=True)
    print("第三项目已换为智慧农场")

    # 4) 删除校园经历整栏（标题 + 一句正文）
    for index in (32, 31):
        element = children[index]
        if element.getparent() is not None:
            element.getparent().remove(element)
    print("已删除校园经历整栏")

    output.parent.mkdir(parents=True, exist_ok=True)
    staged = output.with_name(output.stem + ".staged.docx")
    doc.save(staged)

    # 只替换 word/document.xml，其余部件从基准模板原样复制，保证字节一致
    with zipfile.ZipFile(BASE) as zin, zipfile.ZipFile(staged) as zedit:
        edited_document = zedit.read("word/document.xml")
    with zipfile.ZipFile(BASE) as zin, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = edited_document if item.filename == "word/document.xml" else zin.read(item.filename)
            zout.writestr(item, data)
    staged.unlink()

    with zipfile.ZipFile(output) as archive:
        bad = archive.testzip()
    if bad is not None:
        print(f"ERROR: DOCX 损坏成员：{bad}", file=sys.stderr)
        return 2
    print(f"OK: 已生成 {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
